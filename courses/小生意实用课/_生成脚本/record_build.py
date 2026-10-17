"""Codex 课试讲记录表（给助教填）。python build.py out.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
FONT = "Microsoft YaHei"
DARK, LIME, FILL, EX = "16191F", "C8F031", "FFF9C4", "EEEEEE"
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# 课上的 8 个练习（页码按 70 页版课件）
EXERCISES = [
    ("P28", "打开 Codex、定规矩、做技能"),
    ("P35", "排计划"),
    ("P39", "整理邮箱、起草回复"),
    ("P43", "每日汇总（定时任务）"),
    ("P47", "客户分析"),
    ("P50", "周报"),
    ("P53", "知识库"),
    ("P58", "报价计算器"),
]
EX_NAMES = [f"{p} {n}" for p, n in EXERCISES]
ROWS = 30  # 每张表预留的行数

wb = Workbook()


def header(ws, cols, widths, row=3):
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color=LIME, size=10)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = note
    ws["A2"].font = Font(name=FONT, size=9, color="666666")


def body(ws, ncols, first, last, example=None, prefill=None):
    """example: 第一行示例（灰底斜体）；prefill: {列号: 函数(i)} 预填编号"""
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=FILL)
    if example:
        for c, v in enumerate(example, start=1):
            cell = ws.cell(row=first, column=c, value=v)
            cell.font = Font(name=FONT, size=10, italic=True, color="777777")
            cell.fill = PatternFill("solid", fgColor=EX)
    if prefill:
        start = first + (1 if example else 0)
        for c, f in prefill.items():
            for i, r in enumerate(range(start, last + 1), start=1):
                ws.cell(row=r, column=c, value=f(i))


def dropdown(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    dv.error = "请从下拉选项里选"
    ws.add_data_validation(dv)
    dv.add(rng)


# ── 0. 说明 ──
ws = wb.active
ws.title = "说明"
lines = [
    ("Codex 课试讲记录表", True, 16),
    ("用途：试讲时记下学员在哪里卡住、问了什么、额度在哪用完，课后按记录改课件。", False, 11),
    ("", False, 11),
    ("怎么填", True, 12),
    ("• 黄色格子是要填的；灰色斜体的第一行是示例，照着格式填就行，填完可以删掉。", False, 11),
    ("• 有下拉箭头的格子，直接点选，不用打字。", False, 11),
    ("• 「练习统计」那张表是自动算的，不用填。", False, 11),
    ("", False, 11),
    ("谁填哪张", True, 12),
    ("• 学员情况：开课前登记 + 课上每做完一个练习就更新（助教）", False, 11),
    ("• 卡点记录：谁卡住了就记一条，越具体越好（助教，边走边记）", False, 11),
    ("• 学员提问：学员问的原话（讲师或助教）", False, 11),
    ("• 课后反馈：下课前让学员口头回答或扫码填，助教汇总", False, 11),
    ("• 讲师复盘：讲完当天讲师自己填", False, 11),
    ("", False, 11),
    ("最重要的三件事", True, 12),
    ("1. 免费账号的额度，用到第几个练习就没了 → 填在「学员情况」的「额度用完在」", False, 11),
    ("2. 哪一页大家都卡住 → 「卡点记录」的页码", False, 11),
    ("3. 学员问的原话 → 不要改写，原话最有用", False, 11),
]
for i, (t, b, s) in enumerate(lines, start=1):
    ws.cell(row=i, column=1, value=t).font = Font(name=FONT, bold=b, size=s)
ws.column_dimensions["A"].width = 90

# ── 1. 学员情况 ──
ws = wb.create_sheet("学员情况")
title(ws, "学员情况", "每位学员一行。练习列：每做完一个就选「做完」；没做完选「半途」；没做选「没做」；只看讲师演示选「看演示」。")
cols = ["编号", "做什么行业/工作", "电脑", "ChatGPT 账号", "Codex 装好了吗", "额度用完在"] + EX_NAMES + ["备注"]
widths = [7, 16, 10, 12, 13, 22] + [12] * len(EX_NAMES) + [24]
header(ws, cols, widths)
first, last = 4, 3 + ROWS
example = ["S00", "示例：咖啡店老板", "Windows", "免费", "现场装好", "P47 客户分析",
           "做完", "做完", "做完", "半途", "没做", "没做", "没做", "看演示", "注册卡了 10 分钟"]
body(ws, len(cols), first, last, example, {1: lambda i: f"S{i:02d}"})
dropdown(ws, ["Windows", "Mac", "平板", "其他"], f"C{first}:C{last}")
dropdown(ws, ["免费", "Go", "Plus", "Pro", "公司账号", "没有账号"], f"D{first}:D{last}")
dropdown(ws, ["课前已装", "现场装好", "装不上"], f"E{first}:E{last}")
dropdown(ws, EX_NAMES + ["没用完"], f"F{first}:F{last}")
ex_first_col = 7
dropdown(ws, ["做完", "半途", "没做", "看演示"],
         f"{get_column_letter(ex_first_col)}{first}:{get_column_letter(ex_first_col + len(EX_NAMES) - 1)}{last}")

# ── 2. 卡点记录 ──
ws = wb.create_sheet("卡点记录")
title(ws, "卡点记录", "谁卡住了就记一条。「卡在哪」写具体的，比如「找不到 Codex 指令在哪」，不要只写「不会」。")
cols = ["时间", "课件页码", "哪个练习", "学员编号", "卡点类型", "卡在哪（具体）", "怎么解决的", "花了几分钟", "要改课件吗"]
widths = [8, 9, 22, 9, 16, 34, 28, 9, 10]
header(ws, cols, widths)
first, last = 4, 3 + ROWS * 2
example = ["10:42", "P26", "P28 打开 Codex、定规矩、做技能", "S03", "找不到按钮/入口",
           "找不到「Codex 指令」，以为在聊天框里粘贴", "助教带着点：设置 → 个性化", 5, "要"]
body(ws, len(cols), first, last, example)
dropdown(ws, EX_NAMES + ["不是练习"], f"C{first}:C{last}")
dropdown(ws, ["安装/登录账号", "找不到按钮/入口", "听不懂概念", "不知道怎么跟它说", "AI 结果不对",
              "额度用完", "插件/权限连不上", "网络/电脑问题", "其他"], f"E{first}:E{last}")
dropdown(ws, ["要", "不用"], f"I{first}:I{last}")

# ── 3. 学员提问 ──
ws = wb.create_sheet("学员提问")
title(ws, "学员提问", "记原话，不要改写。最后一列判断：这个问题是不是很多人都会问、要不要写进课件。")
cols = ["时间", "课件页码", "学员编号", "问题原话", "怎么回答的", "几个人问过类似的", "要加进课件吗"]
widths = [8, 9, 9, 40, 34, 10, 11]
header(ws, cols, widths)
first, last = 4, 3 + ROWS
example = ["11:05", "P19", "S07", "那我用 WorkBuddy 可以吗？跟 Codex 有什么不一样？", "可以，方法一样，界面不同", 3, "要"]
body(ws, len(cols), first, last, example)
dropdown(ws, ["要", "不用", "再看看"], f"G{first}:G{last}")

# ── 4. 练习统计（自动） ──
ws = wb.create_sheet("练习统计")
title(ws, "练习统计（自动算，不用填）", "根据「学员情况」和「卡点记录」自动统计。示例行（S00 和第一条示例卡点）也会算进去，正式用前把示例删掉。")
cols = ["练习", "做完", "半途", "没做", "看演示", "额度在这里用完", "卡点次数", "做完比例"]
widths = [30, 8, 8, 8, 9, 14, 10, 10]
header(ws, cols, widths)
s1_first, s1_last = 4, 3 + ROWS
s2_first, s2_last = 4, 3 + ROWS * 2
for i, name in enumerate(EX_NAMES):
    r = 4 + i
    col = get_column_letter(ex_first_col + i)
    rng = f"学员情况!${col}${s1_first}:${col}${s1_last}"
    ws.cell(row=r, column=1, value=name)
    ws.cell(row=r, column=2, value=f'=COUNTIF({rng},"做完")')
    ws.cell(row=r, column=3, value=f'=COUNTIF({rng},"半途")')
    ws.cell(row=r, column=4, value=f'=COUNTIF({rng},"没做")')
    ws.cell(row=r, column=5, value=f'=COUNTIF({rng},"看演示")')
    ws.cell(row=r, column=6, value=f'=COUNTIF(学员情况!$F${s1_first}:$F${s1_last},A{r})')
    ws.cell(row=r, column=7, value=f'=COUNTIF(卡点记录!$C${s2_first}:$C${s2_last},A{r})')
    ws.cell(row=r, column=8, value=f'=IFERROR(B{r}/SUM(B{r}:E{r}),"-")')
    for c in range(1, 9):
        cell = ws.cell(row=r, column=c)
        cell.font = Font(name=FONT, size=10)
        cell.border = BORDER
    ws.cell(row=r, column=8).number_format = "0%"
r = 4 + len(EX_NAMES) + 1
ws.cell(row=r, column=1, value="额度没用完的人数").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=r, column=6, value=f'=COUNTIF(学员情况!$F${s1_first}:$F${s1_last},"没用完")').font = Font(name=FONT, size=10)
ws.cell(row=r + 1, column=1, value="装不上 Codex 的人数").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=r + 1, column=6, value=f'=COUNTIF(学员情况!$E${s1_first}:$E${s1_last},"装不上")').font = Font(name=FONT, size=10)
ws.cell(row=r + 2, column=1, value="要改课件的卡点").font = Font(name=FONT, bold=True, size=10)
ws.cell(row=r + 2, column=6, value=f'=COUNTIF(卡点记录!$I${s2_first}:$I${s2_last},"要")').font = Font(name=FONT, size=10)

# ── 5. 课后反馈 ──
ws = wb.create_sheet("课后反馈")
title(ws, "课后反馈", "下课前问学员这几个问题，每人一行。打分 1–5，5 分最好。")
cols = ["学员编号", "整体打几分（1–5）", "最有用的是哪部分", "哪里听不懂", "回去第一个会用哪个", "会推荐给朋友吗（1–5）", "其他意见"]
widths = [9, 12, 24, 24, 22, 12, 30]
header(ws, cols, widths)
first, last = 4, 3 + ROWS
example = ["S00", 4, "客户分析", "Harness 那页", "P43 每日汇总（定时任务）", 5, "希望多一点时间做练习"]
body(ws, len(cols), first, last, example, {1: lambda i: f"S{i:02d}"})
dropdown(ws, ["1", "2", "3", "4", "5"], f"B{first}:B{last}")
dropdown(ws, ["1", "2", "3", "4", "5"], f"F{first}:F{last}")
dropdown(ws, EX_NAMES + ["还不确定"], f"E{first}:E{last}")

# ── 6. 讲师复盘 ──
ws = wb.create_sheet("讲师复盘")
title(ws, "讲师复盘", "讲完当天填。每个问题写一两句就行。")
qs = [
    "实际讲了多久？哪一部分超时最多？",
    "哪几页讲得太快、学员没跟上？（写页码）",
    "哪几页太啰嗦、可以删？（写页码）",
    "免费账号的学员，额度大概撑到第几个练习？",
    "现场最耗时间的是什么？（装软件、注册、连插件……）",
    "助教够不够？下次要几个？",
    "学员最感兴趣的是哪部分？",
    "下次一定要改的三件事",
]
ws.column_dimensions["A"].width = 44
ws.column_dimensions["B"].width = 70
for i, q in enumerate(qs, start=4):
    a = ws.cell(row=i, column=1, value=q)
    a.font = Font(name=FONT, bold=True, size=10)
    a.alignment = Alignment(vertical="top", wrap_text=True)
    a.border = BORDER
    b = ws.cell(row=i, column=2)
    b.fill = PatternFill("solid", fgColor=FILL)
    b.border = BORDER
    b.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[i].height = 48

wb.save(OUT)
print("saved", OUT)
